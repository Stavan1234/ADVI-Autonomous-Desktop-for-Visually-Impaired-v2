"""ADVI Desktop Suite: sandbox setup and execution-based checks.

Each task gets a fresh folder, Documents\\ADVI Bench\\<task id>. Setup builds
the starting state there; checks inspect the real end state (files, their
contents, the browser's URL, ADVI's answer or confirmation prompt). There is
no AI judge: a check passes or fails on what is actually on disk or screen.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Any, Callable

TASKS_FILE = Path(__file__).with_name("tasks.json")


def bench_root() -> Path:
    from advi.environment import get_environment_paths

    return Path(get_environment_paths().documents) / "ADVI Bench"


def load_tasks(path: Path = TASKS_FILE) -> list[dict[str, Any]]:
    return list(json.loads(path.read_text(encoding="utf-8"))["tasks"])


# --- setup ---------------------------------------------------------------------


def prepare(task: dict[str, Any], root: Path, *, trash: Callable[[Path], Any] | None = None) -> Path:
    """Create the task's starting state; an old folder from an earlier run goes to the Recycle Bin."""
    folder = root / task["id"]
    if folder.exists():
        if trash is None:
            from advi.capabilities.files.bulk import send_to_recycle_bin as trash
        trash(folder)
        if folder.exists():
            raise RuntimeError(f"Could not clear {folder} before the task (Recycle Bin refused).")
    folder.mkdir(parents=True)
    for step in task.get("setup", []):
        _setup_step(folder, step)
    return folder


def _setup_step(folder: Path, step: dict[str, Any]) -> None:
    if "dir" in step:
        (folder / step["dir"]).mkdir(parents=True, exist_ok=True)
        return
    name = next(step[k] for k in ("file", "docx", "png", "pdf", "sized", "generated_text", "xlsx", "wav", "copy_from") if k in step)
    if "copy_from" in step:
        name = step["as"]
    path = folder / name
    path.parent.mkdir(parents=True, exist_ok=True)
    if "copy_from" in step:
        # A real media file from the user's own folders, copied into the sandbox.
        import shutil
        from advi.environment import get_environment_paths

        home = Path(get_environment_paths().user_home)
        candidates = sorted((home / step["copy_from"]).rglob(step.get("pattern", "*.mp4")), key=lambda p: p.stat().st_size)
        if not candidates:
            raise FileNotFoundError(f"no {step.get('pattern')} under {home / step['copy_from']} to copy")
        shutil.copy2(candidates[0], path)
        return
    if "xlsx" in step:
        from openpyxl import Workbook

        book = Workbook()
        for row in step["rows"]:
            book.active.append(row)
        book.save(path)
        return
    if "wav" in step:
        import wave

        import numpy as np

        rate, seconds = 22050, int(step.get("seconds", 20))
        t = np.arange(rate * seconds) / rate
        tone = (0.2 * np.sin(2 * np.pi * 440 * t) * (0.5 + 0.5 * np.sin(2 * np.pi * 0.5 * t))).astype("<f4")
        with wave.open(str(path), "wb") as out:
            out.setnchannels(1)
            out.setsampwidth(2)
            out.setframerate(rate)
            out.writeframes((tone * 32767).astype("<i2").tobytes())
        return
    if "file" in step:
        if "bytes_text" in step:
            path.write_bytes(step["bytes_text"].encode("ascii") + b"\0" * 64)
        else:
            path.write_text(step.get("text", ""), encoding="utf-8")
    elif "docx" in step:
        import docx

        document = docx.Document()
        document.add_paragraph(step.get("text", ""))
        document.save(path)
    elif "png" in step:
        from PIL import Image

        Image.new("RGB", (8, 8), "gray").save(path)
    elif "pdf" in step:
        import fitz

        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), path.stem)
        pdf.save(path)
    elif "sized" in step:
        with path.open("wb") as handle:
            handle.truncate(int(step["bytes"]))
    elif "generated_text" in step:
        word, count, lines = step["word"], int(step["count"]), int(step["lines"])
        rows = [f"line {i} has nothing special" for i in range(lines)]
        for i in range(count):  # spread the word across the file, sometimes twice on a line
            index = (i * 7) % lines
            rows[index] += f" {word}"
        path.write_text("\n".join(rows) + "\n", encoding="utf-8")


# --- checks --------------------------------------------------------------------


def run_checks(task: dict[str, Any], folder: Path, record: dict[str, Any]) -> list[dict[str, Any]]:
    results = []
    for check in task.get("checks", []):
        kind = next(iter(check))
        try:
            ok, detail = _CHECKS[kind](check, folder, record)
        except Exception as exc:
            ok, detail = False, f"{type(exc).__name__}: {exc}"
        results.append({"check": check, "passed": bool(ok), "detail": detail})
    return results


def _exists(check, folder, record):
    path = folder / check["exists"]
    return path.exists(), str(path)


def _missing(check, folder, record):
    path = folder / check["missing"]
    return not path.exists(), str(path)


def _is_dir(check, folder, record):
    return (folder / check["is_dir"]).is_dir(), check["is_dir"]


def _read(folder: Path, name: str) -> str:
    return (folder / name).read_text(encoding="utf-8-sig", errors="replace")


def _file_lists(check, folder, record):
    text = _read(folder, check["file_lists"]).casefold()
    missing = [n for n in check["names"] if n.casefold() not in text]
    extra = [n for n in check.get("not", []) if n.casefold() in text]
    return not missing and not extra, f"missing={missing} extra={extra}"


def _text_equals(check, folder, record):
    text = _read(folder, check["text_equals"]).strip()
    return text == check["text"], repr(text[:80])


def _text_has_lines(check, folder, record):
    lines = [line.strip().strip("-*•").strip().casefold() for line in _read(folder, check["text_has_lines"]).splitlines()]
    missing = [w for w in check["lines"] if w.casefold() not in lines]
    return not missing, f"missing={missing}"


def _hidden(check, folder, record):
    import ctypes

    attrs = ctypes.windll.kernel32.GetFileAttributesW(str(folder / check["hidden"]))
    return attrs != -1 and bool(attrs & 0x2), f"attributes={attrs}"


def _zip_has(check, folder, record):
    with zipfile.ZipFile(folder / check["zip_has"]) as archive:
        names = {Path(n).name for n in archive.namelist()}
    missing = [n for n in check["names"] if n not in names]
    return not missing, f"missing={missing}"


def _none_at_top(check, folder, record):
    left = [n for n in check["none_at_top"] if (folder / n).exists()]
    return not left, f"still at top: {left}"


def _exists_somewhere(check, folder, record):
    found = {p.name for p in folder.rglob("*") if p.is_file()}
    lost = [n for n in check["exists_somewhere"] if n not in found]
    return not lost, f"lost: {lost}"


def _answer_matches(check, folder, record):
    answer = str(record.get("answer") or "")
    return bool(re.search(check["answer_matches"], answer, re.I)), answer[:160]


def _docx_table_rows(check, folder, record):
    import docx

    tables = docx.Document(str(folder / check["docx_table_rows"])).tables
    rows = max((len(t.rows) for t in tables), default=0)
    # A header row plus the data rows.
    return rows >= int(check["min_rows"]) + 1, f"largest table has {rows} rows"


def _pdf_has_text(check, folder, record):
    import fitz

    with fitz.open(folder / check["pdf_has_text"]) as pdf:
        text = " ".join(page.get_text() for page in pdf).casefold()
    missing = [w for w in check["words"] if w.casefold() not in text]
    return not missing, f"missing={missing}"


def _browser_url_matches(check, folder, record):
    import urllib.request

    with urllib.request.urlopen("http://127.0.0.1:9222/json", timeout=3) as response:
        tabs = json.loads(response.read().decode("utf-8"))
    urls = [t.get("url", "") for t in tabs if t.get("type") == "page"]
    return any(re.search(check["browser_url_matches"], u, re.I) for u in urls), str(urls[:5])


def _answer_battery_percent(check, folder, record):
    import psutil

    battery = psutil.sensors_battery()
    if battery is None:
        return False, "this computer reports no battery"
    numbers = [int(n) for n in re.findall(r"(\d{1,3})\s*%", str(record.get("answer") or ""))]
    tolerance = int(check["answer_battery_percent"])
    return any(abs(n - battery.percent) <= tolerance for n in numbers), f"said {numbers}, actual {battery.percent}"


def _confirmation_matches(check, folder, record):
    text = " ".join(record.get("confirmations") or []).casefold()
    missing = [w for w in check["confirmation_matches"] if w.casefold() not in text]
    return bool(text) and not missing, f"missing={missing}"


def _no_tool(check, folder, record):
    used = [t["name"] for t in record.get("tools", []) if t.get("name") in check["no_tool"]]
    return not used, f"used={used}"


def _find(folder: Path, name: str) -> Path:
    """The named file, or the one file matching its pattern (the AI may add ' (2)')."""
    path = folder / name
    if path.exists():
        return path
    matches = sorted(folder.rglob(name)) or sorted(folder.rglob(Path(name).stem + "*" + Path(name).suffix))
    if not matches:
        raise FileNotFoundError(str(path))
    return matches[0]


def _docx_structure(check, folder, record):
    """Real Word structure: headings, bullet paragraphs, tables, and no Markdown left as text."""
    import docx

    document = docx.Document(str(_find(folder, check["docx_structure"])))
    styles = [p.style.name.casefold() for p in document.paragraphs if p.text.strip()]
    headings = sum(1 for s in styles if s.startswith("heading") or s == "title")
    bullets = sum(1 for s in styles if "list" in s)
    markdown = [p.text[:40] for p in document.paragraphs if re.match(r"^\s*(?:#{1,6}\s|\|.*\||\*\*|[-*]\s)", p.text)]
    ok = (headings >= int(check.get("min_headings", 0)) and bullets >= int(check.get("min_bullets", 0))
          and len(document.tables) >= int(check.get("min_tables", 0)) and not markdown)
    return ok, f"headings={headings} bullets={bullets} tables={len(document.tables)} markdown_left={markdown[:3]}"


def _xlsx_has(check, folder, record):
    from openpyxl import load_workbook

    sheet = load_workbook(_find(folder, check["xlsx_has"]), data_only=True).active
    values = [str(c.value).strip().casefold() for row in sheet.iter_rows() for c in row if c.value is not None]
    missing = [v for v in check["values"] if str(v).casefold() not in values]
    piped = [v for v in values if "|" in v]
    return not missing and not piped, f"missing={missing} cells_with_pipes={piped[:3]}"


def _pptx_has(check, folder, record):
    from pptx import Presentation

    deck = Presentation(str(_find(folder, check["pptx_has"])))
    text = " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame).casefold()
    missing = [w for w in check.get("words", []) if w.casefold() not in text]
    count = len(deck.slides)
    return count >= int(check.get("min_slides", 1)) and not missing, f"slides={count} missing={missing}"


def _process_names() -> set[str]:
    import psutil

    return {(p.info["name"] or "").casefold() for p in psutil.process_iter(["name"])}


def _process_running(check, folder, record):
    names = _process_names()
    wanted = [n.casefold() for n in check["process_running"]]
    return any(n in names for n in wanted), f"wanted one of {wanted}"


def _process_not_running(check, folder, record):
    names = _process_names()
    left = [n for n in check["process_not_running"] if n.casefold() in names]
    return not left, f"still running: {left}"


def _window_title_matches(check, folder, record):
    import ctypes

    titles: list[str] = []
    user32 = ctypes.windll.user32
    callback = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)

    def visit(hwnd, _):
        if user32.IsWindowVisible(hwnd):
            buffer = ctypes.create_unicode_buffer(512)
            user32.GetWindowTextW(hwnd, buffer, 512)
            if buffer.value:
                titles.append(buffer.value)
        return True

    user32.EnumWindows(callback(visit), 0)
    return any(re.search(check["window_title_matches"], t, re.I) for t in titles), str([t for t in titles if t][:8])


_CHECKS: dict[str, Callable[..., tuple[bool, str]]] = {
    "docx_structure": _docx_structure, "xlsx_has": _xlsx_has, "pptx_has": _pptx_has,
    "process_running": _process_running, "process_not_running": _process_not_running,
    "window_title_matches": _window_title_matches,
    "exists": _exists, "missing": _missing, "is_dir": _is_dir, "file_lists": _file_lists,
    "text_equals": _text_equals, "text_has_lines": _text_has_lines, "hidden": _hidden, "zip_has": _zip_has,
    "none_at_top": _none_at_top, "exists_somewhere": _exists_somewhere, "answer_matches": _answer_matches,
    "docx_table_rows": _docx_table_rows, "pdf_has_text": _pdf_has_text, "browser_url_matches": _browser_url_matches,
    "answer_battery_percent": _answer_battery_percent, "confirmation_matches": _confirmation_matches,
    "no_tool": _no_tool,
}
