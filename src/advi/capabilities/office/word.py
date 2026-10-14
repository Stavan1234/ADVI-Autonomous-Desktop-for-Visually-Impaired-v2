"""Microsoft Word skill: build a formatted document in Word and save it, in one call.

Typing a document into Word key by key took many model rounds, broke on
focus changes, and stalled on Word's "Product Activation Failed" state. Word's
automation interface keeps working in that state (checked live), so the whole
document is built through it in a private, hidden Word instance:

    markdown-like content -> Word styles (Title, Heading 1-3, bullets,
    numbering, quotes, tables, bold/italic) -> page numbers -> SaveAs
    (.docx) or export (.pdf) -> the saved file is read back -> optionally
    opened in Word for the user, with activation nag dialogs dismissed.

Nothing is ever overwritten: an existing name gets " (2)", " (3)", ...
"""
from __future__ import annotations

import logging
import os
import re
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

logger = logging.getLogger(__name__)

# Built-in style IDs are language independent (style names are localized).
STYLE = {"title": -63, "subtitle": -75, "h1": -2, "h2": -3, "h3": -4, "normal": -1,
         "bullet": -49, "number": -50, "quote": -181}
WD_FORMAT_DOCX = 16
WD_FORMAT_PDF = 17


@dataclass
class Block:
    kind: str                      # title, subtitle, h1..h3, normal, bullet, number, quote, table
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


_INLINE = re.compile(r"(\*\*(.+?)\*\*|__(.+?)__|\*(?!\s)(.+?)(?<!\s)\*)")


def inline_spans(text: str) -> tuple[str, list[tuple[int, int, str]]]:
    """Strip **bold** / *italic* markers; return plain text and (start, end, kind) spans."""
    plain, spans, pos = [], [], 0
    length = 0
    for match in _INLINE.finditer(text):
        before = text[pos:match.start()]
        plain.append(before)
        length += len(before)
        inner = match.group(2) or match.group(3) or match.group(4) or ""
        kind = "italic" if match.group(4) else "bold"
        spans.append((length, length + len(inner), kind))
        plain.append(inner)
        length += len(inner)
        pos = match.end()
    plain.append(text[pos:])
    return "".join(plain), spans


def parse_blocks(content: str, title: str = "") -> list[Block]:
    blocks: list[Block] = []
    if title.strip():
        blocks.append(Block("title", title.strip()))
    table: list[list[str]] = []

    def flush_table() -> None:
        if table:
            blocks.append(Block("table", rows=[row[:] for row in table]))
            table.clear()

    for raw in str(content or "").replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if line.startswith("|") and line.endswith("|") and line.count("|") >= 3:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # |---|---| separator
            table.append(cells)
            continue
        flush_table()
        if not line:
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not blocks:
                blocks.append(Block("title", text))   # a leading "# X" is the document title
            elif level == 1 and blocks and blocks[-1].kind == "title" and blocks[-1].text.casefold() == text.casefold():
                continue                               # title repeated as the first heading
            else:
                blocks.append(Block(f"h{level}", text))
        elif re.match(r"^[-*•]\s+", line):
            blocks.append(Block("bullet", re.sub(r"^[-*•]\s+", "", line)))
        elif re.match(r"^\d+[.)]\s+", line):
            blocks.append(Block("number", re.sub(r"^\d+[.)]\s+", "", line)))
        elif line.startswith(">"):
            blocks.append(Block("quote", line.lstrip("> ").strip()))
        else:
            blocks.append(Block("normal", line))
    flush_table()
    return blocks


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for n in range(2, 1000):
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not candidate.exists():
            return candidate
    raise FileExistsError(f"Too many files named like {path.name}")


def word_available() -> bool:
    if os.name != "nt":
        return False
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"):
            return True
    except OSError:
        return False


def _word_pids() -> set[int]:
    import psutil

    return {p.pid for p in psutil.process_iter(["name"]) if (p.info["name"] or "").upper() == "WINWORD.EXE"}


def build_with_word(blocks: list[Block], target: Path, *, timeout: float = 90.0) -> None:
    """Build and save the document in a private hidden Word; kill it if it hangs."""
    before = _word_pids()
    error: list[BaseException] = []

    def work() -> None:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        app = doc = None
        try:
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            app.ScreenUpdating = False
            doc = app.Documents.Add()
            _fill(doc, blocks)
            if target.suffix.lower() == ".pdf":
                doc.ExportAsFixedFormat(str(target), WD_FORMAT_PDF)
            else:
                doc.SaveAs2(str(target), WD_FORMAT_DOCX)
        except BaseException as exc:  # reported to the caller
            error.append(exc)
        finally:
            try:
                if doc is not None:
                    doc.Close(0)
                if app is not None:
                    app.Quit(0)
            except Exception:
                pass
            pythoncom.CoUninitialize()

    worker = threading.Thread(target=work, name="advi-word", daemon=True)
    worker.start()
    worker.join(timeout)
    if worker.is_alive():
        for pid in _word_pids() - before:   # only the instance this call started
            try:
                import psutil

                psutil.Process(pid).kill()
            except Exception:
                pass
        raise TimeoutError("Word stopped responding while building the document.")
    if error:
        raise error[0]


def _fill(doc: Any, blocks: list[Block]) -> None:
    for block in blocks:
        start = doc.Content.End - 1
        if block.kind == "table":
            _add_table(doc, start, block.rows)
            continue
        plain, spans = inline_spans(block.text)
        rng = doc.Range(start, start)
        rng.InsertAfter(plain)
        rng.Style = doc.Styles(STYLE.get(block.kind, STYLE["normal"]))
        for a, b, kind in spans:
            piece = doc.Range(rng.Start + a, rng.Start + b)
            if kind == "bold":
                piece.Font.Bold = True
            else:
                piece.Font.Italic = True
        rng.InsertParagraphAfter()
    last = doc.Paragraphs(doc.Paragraphs.Count).Range
    if not last.Text.strip() and doc.Paragraphs.Count > 1:
        last.Delete()
    try:
        doc.Sections(1).Footers(1).PageNumbers.Add(1)  # centred page numbers
    except Exception:
        pass


def _add_table(doc: Any, start: int, rows: list[list[str]]) -> None:
    width = max(len(r) for r in rows)
    rng = doc.Range(start, start)
    # The empty paragraph inherits the previous style (e.g. List Number), which
    # numbered every cell in the first live test.
    rng.Style = doc.Styles(STYLE["normal"])
    table = doc.Tables.Add(rng, len(rows), width)
    table.Range.Style = doc.Styles(STYLE["normal"])
    for style in ("Grid Table 4 - Accent 1", "Table Grid"):
        try:
            table.Style = style
            break
        except Exception:
            continue
    for i, row in enumerate(rows, start=1):
        for j in range(width):
            plain, _ = inline_spans(row[j] if j < len(row) else "")
            table.Cell(i, j + 1).Range.Text = plain
    table.Rows(1).Range.Font.Bold = True
    table.AutoFitBehavior(2)  # full page width
    after = doc.Range(doc.Content.End - 1, doc.Content.End - 1)
    after.Style = doc.Styles(STYLE["normal"])
    after.ParagraphFormat.SpaceBefore = 8


def read_back(path: Path) -> dict[str, Any]:
    """Independent check of what was saved."""
    data = path.read_bytes()
    if path.suffix.lower() == ".pdf":
        pages = len(re.findall(rb"/Type\s*/Page[^s]", data))
        return {"ok": data.startswith(b"%PDF") and pages > 0, "pages": pages, "bytes": len(data)}
    import docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return {"ok": bool(text.strip()), "text": text, "words": len(text.split()), "bytes": len(data)}
