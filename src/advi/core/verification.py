from __future__ import annotations

import ctypes
import logging
import os
import re
import urllib.parse
import unicodedata
from pathlib import Path
from typing import Any

from advi.tools.contracts import ToolCall, ToolResult, VerificationStatus
from advi.environment import get_environment_paths

logger = logging.getLogger(__name__)


class ToolVerifier:
    """
    Generalized verification layer for ADVI actions.
    Inspects actual system state to confirm an action succeeded.
    """

    def verify(self, action: ToolCall, result: ToolResult) -> ToolResult:
        """Run verification based on action type.

        IMPORTANT: A verification failure does NOT downgrade a successful execution
        unless there is real evidence of failure (e.g. file not found, wrong URL).
        If there is no specific verifier, or the verifier cannot confirm state,
        the result is returned with verified=False but success unchanged.
        """
        if not result.success:
            # `create_file` is intentionally non-overwriting. When the target
            # already exists, inspect it before deciding whether the request is
            # an idempotent success or a real content conflict.
            if action.action == "create_file" and result.metadata.get("already_exists"):
                try:
                    verified, details = self._verify_create_file(action, result)
                    if verified:
                        return ToolResult(
                            call_id=result.call_id, name=result.name, success=True,
                            output=result.data, human_readable=result.human_readable,
                            error=None, verified=True, verification_status=VerificationStatus.VERIFIED,
                            verification_details={**details, "already_exists": True, "idempotent_match": True},
                            evidence=result.evidence,
                            metadata={**dict(result.metadata), "already_exists": True, "idempotent_match": True},
                            requires_confirmation=result.requires_confirmation,
                            confirmation_reason=result.confirmation_reason,
                            security_flags=result.security_flags,
                        )
                    return ToolResult(
                        call_id=result.call_id, name=result.name, success=False,
                        output=result.data, human_readable=result.human_readable,
                        error=result.error or "The existing file does not match the requested content.",
                        verified=False, verification_status=VerificationStatus.FAILED,
                        verification_details={**details, "already_exists": True, "idempotent_match": False, "definitive_failure": True},
                        evidence=result.evidence, metadata=result.metadata,
                        requires_confirmation=result.requires_confirmation,
                        confirmation_reason=result.confirmation_reason,
                        security_flags=result.security_flags,
                    )
                except Exception as exc:
                    return ToolResult(
                        call_id=result.call_id, name=result.name, success=False, output=result.data,
                        human_readable=result.human_readable, error=result.error, verified=False,
                        verification_status=VerificationStatus.UNCERTAIN,
                        verification_details={"already_exists": True, "definitive_failure": False, "reason": str(exc)},
                        evidence=result.evidence, metadata=result.metadata,
                        requires_confirmation=result.requires_confirmation,
                        confirmation_reason=result.confirmation_reason,
                        security_flags=result.security_flags,
                    )
            return result

        verifier_name = f"_verify_{action.action}"
        verifier = getattr(self, verifier_name, None)

        if verifier is None:
            # No specific verifier: success is preserved, verified stays False.
            return ToolResult(
                call_id=result.call_id,
                name=result.name,
                success=result.success,
                output=result.data,
                human_readable=result.human_readable,
                error=result.error,
                verified=result.verified,
                verification_status=(
                    VerificationStatus.VERIFIED if result.verified else VerificationStatus.UNCERTAIN
                ),
                verification_details={"method": "no_verifier"},
                evidence=result.evidence,
                metadata=result.metadata,
                requires_confirmation=result.requires_confirmation,
                confirmation_reason=result.confirmation_reason,
                security_flags=result.security_flags,
            )

        try:
            verified, details = verifier(action, result)
            return ToolResult(
                call_id=result.call_id,
                name=result.name,
                # Only downgrade success if verification found POSITIVE evidence of failure
                success=result.success if verified else (
                    False if details.get("definitive_failure") else result.success
                ),
                output=result.data,
                human_readable=result.human_readable,
                error=(
                    result.error
                    if verified
                    else (
                        f"Verification failed: {details.get('reason', 'Condition not met')}"
                        if details.get("definitive_failure")
                        else result.error
                    )
                ),
                verified=verified,
                verification_status=(VerificationStatus.VERIFIED if verified else (
                    VerificationStatus.FAILED if details.get("definitive_failure") else VerificationStatus.UNCERTAIN
                )),
                verification_details=details,
                evidence=result.evidence,
                metadata=result.metadata,
                requires_confirmation=result.requires_confirmation,
                confirmation_reason=result.confirmation_reason,
                security_flags=result.security_flags,
            )
        except Exception as exc:
            logger.warning("Verification exception for %s: %s", action.action, exc)
            return ToolResult(
                call_id=result.call_id,
                name=result.name,
                success=result.success,   # preserve success on verifier crash
                output=result.data,
                human_readable=result.human_readable,
                error=result.error,
                verified=False,
                verification_status=VerificationStatus.UNCERTAIN,
                verification_details={"error": str(exc)},
                evidence=result.evidence,
                metadata=result.metadata,
                requires_confirmation=result.requires_confirmation,
                confirmation_reason=result.confirmation_reason,
                security_flags=result.security_flags,
            )

    @staticmethod
    def summarize(results: list[ToolResult]) -> dict[str, Any]:
        """Summarize evidence across meaningful action results."""
        meaningful = list(results)
        if not meaningful:
            return {"status": "verified", "verified": 0, "uncertain": 0, "failed": 0}
        failed = [r for r in meaningful if not r.success or r.verification_status == VerificationStatus.FAILED]
        uncertain = [r for r in meaningful if r.success and r.verification_status == VerificationStatus.UNCERTAIN]
        verified = [r for r in meaningful if r.success and r.verification_status == VerificationStatus.VERIFIED]
        status = "failed" if failed else ("uncertain" if uncertain else "verified")
        return {"status": status, "verified": len(verified), "uncertain": len(uncertain), "failed": len(failed)}

    # -------------------------------------------------------------
    # File Operations Verification
    # -------------------------------------------------------------
    def _verify_save_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        """Verify file was saved by checking the real filesystem path."""
        # Use the resolved path stored by FileCapability in metadata, or fall back to parameter
        path_str = (
            result.metadata.get("resolved_path")
            or action.parameters.get("path")
            or action.parameters.get("filepath")
            or result.data
        )
        if not path_str:
            return True, {"reason": "no_path_specified"}

        # Resolve Desktop-relative paths the same way FileCapability does
        p = Path(str(path_str))
        if not p.is_absolute():
            parts = p.parts
            if parts and parts[0].lower() == "desktop":
                desktop = get_environment_paths().desktop
                p = desktop / Path(*parts[1:])
            else:
                desktop = get_environment_paths().desktop
                p = desktop / p

        p = p.resolve()
        exists = p.exists()
        size = p.stat().st_size if exists else 0
        if not exists:
            return False, {
                "path": str(p),
                "exists": False,
                "size_bytes": 0,
                "definitive_failure": True,
                "reason": "file_missing",
            }
        expected_content = action.parameters.get("content")
        content_matches = None
        if expected_content is not None:
            expected_text = str(expected_content)
            try:
                suffix = p.suffix.casefold()
                if suffix in {".pdf", ".docx", ".xlsx", ".pptx"}:
                    expected_pages = action.parameters.get("pages")
                    expected_slides = action.parameters.get("slides")
                    # Structured documents are binary containers. Verify that the
                    # native parser can reopen them and that the requested text is
                    # represented in the parsed content; never byte/text-compare
                    # the package itself.
                    visual_sanity: dict[str, Any] | None = None
                    style_requested = bool(
                        action.parameters.get("style")
                        or action.parameters.get("styling")
                        or action.parameters.get("design")
                        or action.parameters.get("fonts")
                        or action.parameters.get("hierarchy")
                        or action.parameters.get("tables")
                    )
                    if suffix == ".pdf":
                        import fitz  # PyMuPDF
                        document = fitz.open(str(p))
                        try:
                            actual_content = "\n".join(str(page.get_text("text") or "") for page in document)
                            structure_ok = len(document) > 0
                            if expected_pages not in (None, ""):
                                structure_ok = structure_ok and len(document) == max(1, int(expected_pages))

                            # A binary PDF existing on disk is not enough when the
                            # user asked for a designed document. Perform a cheap
                            # render sanity check: detect pages that are genuinely
                            # empty rather than attempting to judge aesthetics.
                            checked_pages: list[dict[str, Any]] = []
                            for page_index, page in enumerate(document):
                                text_len = len((page.get_text("text") or "").strip())
                                image_count = len(page.get_images(full=True))
                                drawing_count = len(page.get_drawings())
                                pix = page.get_pixmap(matrix=fitz.Matrix(0.30, 0.30), alpha=False)
                                samples = pix.samples
                                dark_pixels = 0
                                total_pixels = max(1, int(pix.width * pix.height))
                                # RGB triplets; count pixels with meaningful ink.
                                for offset in range(0, len(samples) - 2, 3):
                                    if samples[offset] < 245 or samples[offset + 1] < 245 or samples[offset + 2] < 245:
                                        dark_pixels += 1
                                ink_ratio = dark_pixels / total_pixels
                                page_valid = bool(text_len >= 12 or image_count or drawing_count or ink_ratio >= 0.001)
                                checked_pages.append({
                                    "page": page_index + 1,
                                    "text_chars": text_len,
                                    "images": image_count,
                                    "drawings": drawing_count,
                                    "ink_ratio": round(ink_ratio, 5),
                                    "valid": page_valid,
                                })
                            visual_sanity = {
                                "checked_pages": checked_pages,
                                "checked_page_count": len(checked_pages),
                                "document_page_count": len(document),
                                "valid": bool(checked_pages) and len(checked_pages) == len(document) and all(item["valid"] for item in checked_pages),
                                "method": "pdf_low_resolution_render_sanity",
                            }
                        finally:
                            document.close()
                    elif suffix == ".docx":
                        from docx import Document
                        document = Document(str(p))
                        actual_content = "\n".join([p.text for p in document.paragraphs] + [
                            " | ".join(cell.text for cell in row.cells)
                            for table in document.tables for row in table.rows
                        ])
                        structure_ok = True
                    elif suffix == ".xlsx":
                        from openpyxl import load_workbook
                        workbook = load_workbook(filename=str(p), read_only=True, data_only=True)
                        try:
                            parts = []
                            for sheet in workbook.worksheets:
                                for row in sheet.iter_rows(values_only=True):
                                    values = ["" if value is None else str(value) for value in row]
                                    if any(values):
                                        parts.append("\t".join(values))
                            actual_content = "\n".join(parts)
                            structure_ok = True
                        finally:
                            workbook.close()
                    else:  # .pptx
                        from pptx import Presentation
                        presentation = Presentation(str(p))
                        parts = []
                        for slide in presentation.slides:
                            for shape in slide.shapes:
                                shape_text = getattr(shape, "text", "")
                                if shape_text and shape_text.strip():
                                    parts.append(shape_text.strip())
                        actual_content = "\n".join(parts)
                        structure_ok = True
                        if expected_slides not in (None, ""):
                            structure_ok = len(presentation.slides) == max(1, int(expected_slides))
                    def normalize(value: Any) -> str:
                        # Native document parsers may normalize narrow/no-break
                        # spaces differently from the original source text.
                        normalized = unicodedata.normalize("NFKC", str(value or ""))
                        # PDF/Office text extraction may substitute unsupported
                        # punctuation with placeholder characters. For structured
                        # content validation, compare the meaningful word/number
                        # sequence rather than fragile glyph-level punctuation.
                        normalized = re.sub(r"[^\w]+", " ", normalized, flags=re.UNICODE)
                        return re.sub(r"\s+", " ", normalized).strip().casefold()
                    expected_norm = normalize(expected_text)
                    actual_norm = normalize(actual_content)
                    if not structure_ok or not expected_norm:
                        content_matches = False
                    elif expected_norm in actual_norm:
                        content_matches = True
                    else:
                        # Renderers for PDF/Office documents may insert line breaks,
                        # headers, page artifacts, or generated layout text. Exact
                        # whole-document containment is therefore too brittle. Use a
                        # few meaningful anchors from the requested source instead.
                        fragments = [
                            normalize(part)
                            for part in re.split(r"\n{1,2}", expected_text)
                            if len(normalize(part)) >= 18
                        ]
                        if len(fragments) < 2 and len(expected_norm) >= 60:
                            fragments = [expected_norm[i:i + 90] for i in range(0, min(len(expected_norm), 360), 90)]
                        unique_fragments: list[str] = []
                        for fragment in fragments:
                            if fragment and fragment not in unique_fragments:
                                unique_fragments.append(fragment)
                        anchors = unique_fragments[:5]
                        matched = sum(1 for fragment in anchors if fragment in actual_norm)
                        threshold = 1 if len(anchors) <= 2 else 3
                        content_matches = matched >= threshold
                else:
                    actual_content = p.read_text(encoding="utf-8", errors="replace")
                    content_matches = actual_content == expected_text
            except (OSError, ValueError, RuntimeError) as exc:
                return False, {
                    "path": str(p),
                    "exists": True,
                    "size_bytes": size,
                    "definitive_failure": False,
                    "reason": f"content_read_failed: {exc}",
                }
            except ImportError as exc:
                return False, {
                    "path": str(p),
                    "exists": True,
                    "size_bytes": size,
                    "definitive_failure": False,
                    "reason": f"content_reader_unavailable: {exc}",
                }
        structured = p.suffix.casefold() in {".pdf", ".docx", ".xlsx", ".pptx"}
        # Structure validates that a binary container can be reopened, while
        # content validation proves it represents the user's requested payload.
        # A known content mismatch is therefore a definitive verification failure,
        # even for a structurally valid PDF/Office document.
        if structured:
            if content_matches is False:
                verified = False
                definitive_failure = True
                reason = "content_mismatch"
            else:
                verified = True
                definitive_failure = False
                reason = None
        else:
            verified = content_matches if content_matches is not None else True
            definitive_failure = content_matches is False
            reason = "content_mismatch" if content_matches is False else None
        details: dict[str, Any] = {
            "path": str(p),
            "exists": True,
            "size_bytes": size,
            "content_matches": content_matches,
            "verification_mode": "native_structure_and_content" if structured else "exact",
            "definitive_failure": definitive_failure,
            "reason": reason,
        }
        if structured and 'style_requested' in locals():
            details["style_requested"] = style_requested
            if visual_sanity is not None:
                details["visual_sanity"] = visual_sanity
        return bool(verified), details

    def _verify_create_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_save_file(action, result)


    def _verify_create_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path_str = result.metadata.get("resolved_path") or result.data or action.parameters.get("path")
        if not path_str:
            return False, {"definitive_failure": False, "reason": "missing_folder_path"}
        p = Path(str(path_str)).resolve()
        exists = p.is_dir()
        return exists, {"path": str(p), "exists": exists, "definitive_failure": not exists, "reason": None if exists else "folder_missing"}

    def _verify_open_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path_str = result.metadata.get("resolved_path") or result.data or action.parameters.get("path")
        exists = bool(path_str and Path(str(path_str)).is_file())
        requested = bool(result.metadata.get("open_requested"))
        # The OS association API confirms the open request, not the foreground UI.
        # Keep this verifier explicit so live Windows validation can later add stronger UI evidence.
        return requested and exists, {"path": str(path_str) if path_str else None, "exists": exists, "open_requested": requested, "method": "os_association_request"}

    def _verify_open_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path_str = result.metadata.get("resolved_path") or result.data or action.parameters.get("path")
        exists = bool(path_str and Path(str(path_str)).is_dir())
        requested = bool(result.metadata.get("open_requested"))
        return requested and exists, {"path": str(path_str) if path_str else None, "exists": exists, "open_requested": requested, "method": "os_association_request"}

    def _verify_find_application(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        if isinstance(data, dict):
            valid = bool(data.get("name"))
            return valid, {"method":"windows_application_discovery", "found":valid, "name":data.get("name"), "source":data.get("source")}
        if isinstance(data, list) and not data:
            return True, {"method":"windows_application_discovery", "found":False, "query":action.parameters.get("name") or action.parameters.get("query")}
        return False, {"method":"windows_application_discovery", "found":False, "definitive_failure":True, "reason":"invalid_application_discovery_result"}

    def _verify_list_applications(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        rows = result.data
        valid = isinstance(rows, list) and all(isinstance(row, dict) and "name" in row and "running" in row for row in rows)
        return valid, {"method": "watchtower_catalog_snapshot", "valid_snapshot": valid, "count": len(rows) if isinstance(rows, list) else 0}

    def _verify_append_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path_str = result.metadata.get("resolved_path") or action.parameters.get("path") or result.data
        if not path_str:
            return False, {"definitive_failure": False, "reason": "no_path_specified"}
        p = Path(str(path_str))
        if not p.is_absolute():
            desktop = get_environment_paths().desktop
            p = desktop / Path(*p.parts[1:]) if p.parts and p.parts[0].lower() == "desktop" else desktop / p
        p = p.resolve()
        if not p.exists():
            return False, {"path": str(p), "exists": False, "definitive_failure": True, "reason": "file_missing"}
        content = str(action.parameters.get("content") or action.parameters.get("value") or "")
        text = p.read_text(encoding="utf-8", errors="replace")
        ok = text.endswith(content) if content else True
        return ok, {"path": str(p), "exists": True, "contains_appended_content": ok, "definitive_failure": bool(content and not ok)}

    def _verify_replace_file_text(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path_str = result.metadata.get("resolved_path") or action.parameters.get("path") or result.data
        old_text = str(action.parameters.get("old_text") or "")
        new_text = str(action.parameters.get("new_text") or "")
        if not path_str:
            return False, {"definitive_failure": False, "reason": "no_path_specified"}
        p = Path(str(path_str)).resolve()
        if not p.exists():
            return False, {"path": str(p), "exists": False, "definitive_failure": True, "reason": "file_missing"}
        text = p.read_text(encoding="utf-8", errors="replace")
        old_after = text.count(old_text) if old_text else 0
        replacement_present = new_text in text if new_text else True
        occurrences_before = result.metadata.get("occurrences_before")
        expected_old_after = max(0, int(occurrences_before) - 1) if isinstance(occurrences_before, int) else None
        if expected_old_after is not None:
            ok = replacement_present and old_after == expected_old_after
        else:
            ok = replacement_present and (old_after == 0)
        return ok, {
            "path": str(p),
            "replacement_present": replacement_present,
            "old_text_count_after": old_after,
            "occurrences_before": occurrences_before,
            "expected_old_text_count_after": expected_old_after,
            "definitive_failure": not ok,
            "reason": "replacement_state_mismatch" if not ok else None,
        }

    def _verify_read_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        path = result.metadata.get("resolved_path") or result.metadata.get("path")
        readable = isinstance(data, (str, bytes))
        return readable, {
            "path": path,
            "readable": readable,
            "length": len(data) if hasattr(data, "__len__") else None,
            "source": "executor_payload",
        }

    def _verify_search_file_content(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        output = result.data
        valid = isinstance(output, list)
        details = {
            "query": action.parameters.get("query") or action.parameters.get("text"),
            "valid_search_result": valid,
            "count": len(output) if valid else None,
            "files_scanned": result.metadata.get("files_scanned"),
        }
        if not valid:
            details["definitive_failure"] = True
            details["reason"] = "invalid_content_search_result"
        return valid, details

    def _verify_find_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        output = result.data
        valid = isinstance(output, list)
        timed_out = bool(result.metadata.get("timed_out") or result.verification_details.get("timed_out"))
        details = {
            "method": "filesystem_scan",
            "query": action.parameters.get("name") or action.parameters.get("filename") or action.parameters.get("query"),
            "scope": result.metadata.get("scope"),
            "count": len(output) if valid else None,
            "timed_out": timed_out,
            "complete_search": not timed_out,
            "valid_search_result": valid,
        }
        if not valid:
            details.update({"definitive_failure": True, "reason": "invalid_file_search_result"})
            return False, details
        if timed_out:
            details.update({"definitive_failure": False, "reason": "search_time_limit_reached"})
            return False, details
        return True, details

    def _verify_list_files(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        # FileCapability returns a structured listing payload:
        # {directory, files: [...], folders: [...], ...}. Older verification
        # code assumed the raw result was a list, so perfectly valid listings
        # were marked uncertain and poisoned the task. Accept both shapes.
        payload = result.data
        if isinstance(payload, list):
            directory = result.metadata.get("directory")
            count = len(payload)
            valid = True
        elif isinstance(payload, dict):
            files = payload.get("files")
            folders = payload.get("folders")
            valid = isinstance(files, list) and (folders is None or isinstance(folders, list))
            directory = payload.get("directory") or result.metadata.get("directory")
            count = (len(files) + len(folders or [])) if valid else None
        else:
            directory = result.metadata.get("directory")
            count = None
            valid = False
        return valid, {
            "directory": directory,
            "valid_listing": valid,
            "count": count,
        }

    def _verify_wait(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        duration = action.parameters.get("duration") or action.parameters.get("seconds")
        return True, {"method": "elapsed_wait", "duration": duration}

    def _verify_list_windows(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        rows = result.data
        valid = isinstance(rows, list) and all(isinstance(row, dict) and "hwnd" in row and "title" in row for row in rows)
        return valid, {"method": "windows_enumwindows", "valid_snapshot": valid, "count": len(rows) if valid else 0}

    def _verify_close_window(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        hwnd = result.metadata.get("hwnd") or action.parameters.get("hwnd")
        if os.name != "nt":
            # The native desktop capability is Windows-only; preserve an already
            # verified adapter result on other hosts without pretending to inspect HWNDs.
            return bool(result.verified), {"method": "executor_verification", "verified_by_executor": bool(result.verified)}
        if not hwnd:
            return False, {"definitive_failure": False, "reason": "missing_window_handle"}
        try:
            exists = bool(ctypes.windll.user32.IsWindow(int(hwnd)))
        except Exception as exc:
            return False, {"definitive_failure": False, "reason": f"window_probe_failed: {exc}", "hwnd": hwnd}
        return not exists, {"method": "windows_window_absence", "hwnd": int(hwnd), "exists_after_close": exists, "definitive_failure": exists, "reason": "window_still_exists" if exists else None}

    def _verify_inspect_window(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        valid = isinstance(data, dict) and bool(data.get("hwnd"))
        return valid, {
            "method": "window_inspection_payload",
            "hwnd": data.get("hwnd") if isinstance(data, dict) else None,
            "control_count": data.get("control_count") if isinstance(data, dict) else None,
            "valid_snapshot": valid,
            "definitive_failure": not valid,
            "reason": None if valid else "invalid_window_inspection",
        }

    def _verify_inspect_web_page(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        valid = isinstance(data, list)
        return valid, {
            "method": "dom_interactive_inventory",
            "count": len(data) if isinstance(data, list) else 0,
            "tab_id": result.metadata.get("tab_id"),
            "definitive_failure": not valid,
        }

    def _verify_type_web_element(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        before = result.metadata.get("before")
        after = result.metadata.get("after")
        expected = result.metadata.get("expected")
        valid = isinstance(after, str) and isinstance(expected, str) and (after == expected or expected in after)
        return valid, {
            "method": "dom_input_value",
            "selector": result.metadata.get("selector"),
            "before": before,
            "after": after,
            "expected": expected,
            "definitive_failure": not valid,
            "reason": None if valid else "input_value_mismatch",
        }

    def _verify_read_web_page(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        text = result.data
        valid = isinstance(text, str)
        return valid, {
            "tab_id": result.metadata.get("tab_id"),
            "content_available": valid,
            "length": len(text) if valid else None,
        }

    def _verify_type_text(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        observation = result.metadata.get("observation") or {}
        observed = observation.get("text_delta_observed")
        if observed is True:
            return True, {"method": "uia_control_value_delta", "before_text": observation.get("before_text"), "after_text": observation.get("after_text"), "text_delta_observed": True}
        return False, {"definitive_failure": False, "method": observation.get("method") or "none", "reason": "text_effect_not_observed", "before_text": observation.get("before_text"), "after_text": observation.get("after_text")}

    def _verify_click(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        observation = result.metadata.get("observation") or {}
        if observation.get("state_changed") is True:
            return True, {
                "method": observation.get("method") or "uia_state_delta",
                "target": result.metadata.get("target"),
                "state_changed": True,
            }
        if observation.get("target_focus_changed") is True:
            return True, {
                "method": "uia_focus_transition",
                "target": result.metadata.get("target"),
                "target_focus_changed": True,
            }
        return False, {
            "definitive_failure": False,
            "method": observation.get("method") or "none",
            "reason": "click_effect_not_definitively_observed",
        }

    def _verify_double_click(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_click(action, result)

    def _verify_right_click(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_click(action, result)

    def _verify_close_document(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.metadata.get("verification_details") or result.metadata or {})
        if details.get("close_dialog"):
            return False, {**details, "definitive_failure": False, "reason": "close_dialog_requires_choice", "method": "windows_dialog_observation"}
        closed = bool(details.get("closed"))
        return closed, {
            **details,
            "method": details.get("method") or "ctrl_w_window_observation",
            "definitive_failure": not closed,
            "reason": None if closed else "document_remains_open",
        }

    def _verify_save_document_as(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.metadata.get("verification_details") or result.metadata or {})
        path = details.get("resolved_path") or details.get("path") or action.parameters.get("path")
        exists = bool(path and Path(str(path)).is_file())
        dirty = details.get("dirty")
        verified = bool(exists and dirty is False)
        return verified, {
            **details,
            "method": "saved_path_and_clean_window_state",
            "path_exists": exists,
            "definitive_failure": bool(path and not exists),
            "reason": None if verified else "save_as_postcondition_not_observed",
        }

    def _verify_save_document(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.metadata.get("verification_details") or result.metadata or {})
        path = details.get("document_path")
        after_dirty = details.get("after_dirty")
        dialog = details.get("save_dialog")
        if dialog:
            return False, {**details, "definitive_failure": False, "reason": "save_dialog_requires_choice", "method": "windows_dialog_observation"}
        exists = bool(path and Path(str(path)).is_file())
        verified = bool(path and exists and after_dirty is False)
        return verified, {
            **details,
            "method": "document_path_and_dirty_state",
            "path_exists": exists,
            "definitive_failure": bool(path and not exists),
            "reason": None if verified else "save_postcondition_not_observed",
        }

    # -------------------------------------------------------------
    # Desktop Application Verification
    # -------------------------------------------------------------
    def _verify_new_document(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.metadata.get("verification_details") or result.metadata or {})
        focused = bool(details.get("focused"))
        changed = bool(details.get("title_changed"))
        new_hwnd = details.get("after_hwnd")
        old_hwnd = details.get("before_hwnd")
        created_identity = bool(new_hwnd and old_hwnd and int(new_hwnd) != int(old_hwnd)) if new_hwnd and old_hwnd else False
        if not focused:
            return False, {**details, "definitive_failure": True, "reason": "target_window_not_focused"}
        if changed or created_identity:
            return True, {**details, "method": "document_window_or_title_change"}
        return False, {**details, "definitive_failure": False, "reason": "new_document_postcondition_not_observed", "method": "focus_only"}

    def _verify_open_application(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        hwnd = result.metadata.get("hwnd") or (result.data if isinstance(result.data, int) else None)
        if not hwnd:
            return False, {
                "status": "no_hwnd",
                "definitive_failure": True,
                "reason": "No window handle was captured for the launched application.",
            }

        user32 = ctypes.windll.user32
        is_window = bool(user32.IsWindow(hwnd))
        is_visible = bool(user32.IsWindowVisible(hwnd))
        actual_pid = ctypes.c_ulong()
        user32.GetWindowThreadProcessId(int(hwnd), ctypes.byref(actual_pid))
        expected_pid = result.metadata.get("pid")

        process_match = None
        if expected_pid:
            process_match = int(actual_pid.value) == int(expected_pid)
        application = str(action.parameters.get("application") or result.metadata.get("application") or "").strip()
        process_name = None
        try:
            import psutil
            if actual_pid.value:
                process_name = psutil.Process(int(actual_pid.value)).name()
        except Exception:
            pass

        # Windows application launchers can return a wrapper PID while the
        # user-facing window belongs to a different process (Store apps,
        # ApplicationFrameHost, packaged apps, or shell activation). A PID
        # mismatch alone is therefore not a sufficient identity failure.
        title = ""
        try:
            buffer = ctypes.create_unicode_buffer(1024)
            user32.GetWindowTextW(int(hwnd), buffer, len(buffer))
            title = buffer.value or ""
        except Exception:
            pass

        normalize = lambda value: re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())
        app_norm = normalize(application)
        title_norm = normalize(title)
        process_norm = normalize(Path(str(process_name or "")).stem)

        process_aliases = {
            "notepad": {"notepad", "applicationframehost"},
            "calculator": {"calculatorapp", "applicationframehost", "calculator"},
            "fileexplorer": {"explorer"},
            "explorer": {"explorer"},
            "googlechrome": {"chrome"},
            "chrome": {"chrome"},
            "microsoftedge": {"msedge"},
            "edge": {"msedge"},
            "microsoftword": {"winword"},
            "word": {"winword"},
            "microsoftexcel": {"excel"},
            "excel": {"excel"},
            "microsoftpowerpoint": {"powerpnt"},
            "powerpoint": {"powerpnt"},
        }
        expected_processes = process_aliases.get(app_norm, set())
        semantic_process_match = process_norm in expected_processes if expected_processes else False
        title_match = bool(app_norm and (app_norm in title_norm or title_norm.startswith(app_norm)))
        identity_match = bool(process_match is True or semantic_process_match or title_match)

        valid = is_window and is_visible and identity_match
        details = {
            "hwnd": int(hwnd),
            "is_window": is_window,
            "is_visible": is_visible,
            "actual_pid": int(actual_pid.value),
            "expected_pid": int(expected_pid) if expected_pid else None,
            "process_match": process_match,
            "semantic_process_match": semantic_process_match,
            "title_match": title_match,
            "identity_match": identity_match,
            "identity_source": ("pid" if process_match is True else "process_name" if semantic_process_match else "window_title" if title_match else None),
            "process_name": process_name,
            "window_title": title,
            "application": application,
        }
        return valid, details

    def _verify_focus_window(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        user32 = ctypes.windll.user32
        foreground_hwnd = user32.GetForegroundWindow()
        expected_hwnd = result.metadata.get("hwnd")
        if expected_hwnd:
            is_foreground = (foreground_hwnd == expected_hwnd)
            return is_foreground, {
                "expected_hwnd": expected_hwnd,
                "foreground_hwnd": foreground_hwnd,
                "is_foreground": is_foreground,
            }
        return bool(foreground_hwnd), {"foreground_hwnd": foreground_hwnd}


    # -------------------------------------------------------------
    # File lifecycle verification
    # -------------------------------------------------------------
    def _verify_delete_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path = result.metadata.get("resolved_path") or result.metadata.get("path") or result.data or action.parameters.get("path")
        if not path:
            return False, {"definitive_failure": False, "reason": "missing_path"}
        exists = Path(str(path)).exists()
        return not exists, {"method": "filesystem_absence", "path": str(path), "exists": exists, "definitive_failure": exists, "reason": "path_still_exists" if exists else None}

    def _verify_copy_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        source = result.metadata.get("source") or action.parameters.get("source") or action.parameters.get("from") or action.parameters.get("path")
        destination = result.metadata.get("destination") or action.parameters.get("destination") or action.parameters.get("to") or action.parameters.get("target")
        if not destination:
            return False, {"definitive_failure": False, "reason": "missing_destination"}
        destination_exists = Path(str(destination)).is_file()
        source_exists = Path(str(source)).is_file() if source else None
        return destination_exists, {"method": "filesystem_copy_postcondition", "source": str(source) if source else None, "destination": str(destination), "source_exists": source_exists, "destination_exists": destination_exists, "definitive_failure": not destination_exists, "reason": "destination_missing" if not destination_exists else None}

    def _verify_move_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        source = result.metadata.get("source") or action.parameters.get("source") or action.parameters.get("from") or action.parameters.get("path")
        destination = result.metadata.get("destination") or action.parameters.get("destination") or action.parameters.get("to") or action.parameters.get("target")
        if not source or not destination:
            return False, {"definitive_failure": False, "reason": "missing_move_paths"}
        source_exists = Path(str(source)).exists()
        destination_exists = Path(str(destination)).is_file()
        ok = destination_exists and not source_exists
        return ok, {"method": "filesystem_move_postcondition", "source": str(source), "destination": str(destination), "source_exists": source_exists, "destination_exists": destination_exists, "definitive_failure": not ok, "reason": "move_postcondition_failed" if not ok else None}

    def _verify_rename_file(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_move_file(action, result)

    def _verify_copy_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        source = Path(str(result.metadata.get("source") or action.parameters.get("source") or action.parameters.get("path") or ""))
        destination = Path(str(result.metadata.get("destination") or result.data or action.parameters.get("destination") or ""))
        if not source or not destination:
            return False, {"definitive_failure": False, "reason": "folder_paths_unavailable"}
        ok = source.is_dir() and destination.is_dir()
        return ok, {"source": str(source), "destination": str(destination), "source_exists": source.is_dir(), "destination_exists": destination.is_dir(), "definitive_failure": not ok, "reason": None if ok else "copy_not_verified"}

    def _verify_move_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        source = Path(str(result.metadata.get("source") or action.parameters.get("source") or action.parameters.get("path") or ""))
        destination = Path(str(result.metadata.get("destination") or result.data or action.parameters.get("destination") or ""))
        if not source or not destination:
            return False, {"definitive_failure": False, "reason": "folder_paths_unavailable"}
        ok = (not source.exists()) and destination.is_dir()
        return ok, {"source": str(source), "destination": str(destination), "source_exists": source.exists(), "destination_exists": destination.is_dir(), "definitive_failure": not ok, "reason": None if ok else "move_not_verified"}

    def _verify_rename_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        source = Path(str(result.metadata.get("source") or action.parameters.get("path") or action.parameters.get("source") or ""))
        destination = Path(str(result.metadata.get("destination") or result.data or ""))
        if not source or not destination:
            return False, {"definitive_failure": False, "reason": "folder_paths_unavailable"}
        ok = (not source.exists()) and destination.is_dir()
        return ok, {"source": str(source), "destination": str(destination), "source_exists": source.exists(), "destination_exists": destination.is_dir(), "definitive_failure": not ok, "reason": None if ok else "rename_not_verified"}

    def _verify_delete_folder(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        path = result.metadata.get("path") or result.metadata.get("resolved_path") or result.data or action.parameters.get("path")
        if not path:
            return False, {"definitive_failure": False, "reason": "missing_folder_path"}
        exists = Path(str(path)).exists()
        return not exists, {"method": "filesystem_directory_absence", "path": str(path), "exists": exists, "definitive_failure": exists, "reason": "directory_still_exists" if exists else None}

    # -------------------------------------------------------------
    # Input/interaction verification
    # -------------------------------------------------------------
    def _verify_press_key(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        observation = result.metadata.get("observation") or {}
        if observation.get("state_changed") is True or observation.get("event_dispatched") is True:
            return True, {"method": observation.get("method") or "input_observation", **observation}
        # A keypress such as Escape/Tab can legitimately leave no inspectable state
        # change. Preserve successful transport without falsely asserting a UI change.
        return False, {"definitive_failure": False, "method": "input_dispatch", "reason": "postcondition_not_observable"}

    def _verify_hotkey(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_press_key(action, result)

    def _verify_scroll(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        observation = result.metadata.get("observation") or {}
        if observation.get("state_changed") is True or observation.get("scroll_position_changed") is True:
            return True, {"method": observation.get("method") or "scroll_observation", **observation}
        # Scroll can be valid at a page boundary with no observable position delta.
        return False, {"definitive_failure": False, "method": "scroll_dispatch", "reason": "scroll_postcondition_not_observable"}

    def _verify_read_screen(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        text = result.data
        valid = isinstance(text, str) and bool(text.strip())
        return valid, {"method": "tesseract_ocr", "characters": len(text) if isinstance(text, str) else 0, "text_observed": valid, "definitive_failure": not valid, "reason": "no_text_recognized" if not valid else None}

    # -------------------------------------------------------------
    # MCP verification
    # -------------------------------------------------------------
    def _verify_mcp_list_tools(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        count = result.metadata.get("tool_count")
        valid = bool(result.success and isinstance(result.data, list))
        return valid, {"method": "mcp_tool_discovery", "tool_count": count, "server": result.metadata.get("server")}

    def _verify_mcp_call(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        valid = bool(result.success and result.verified)
        return valid, {"method": "mcp_call_result", "server": result.metadata.get("server"), "tool": result.metadata.get("tool")}

    # -------------------------------------------------------------
    # Browser interaction verification
    # -------------------------------------------------------------
    def _verify_click_web_element(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.metadata.get("verification_details") or result.verification_details or {})
        metadata = dict(result.metadata or {})
        playback = details.get("playback") or details.get("player") or {}
        if (
            details.get("playing") is True
            or metadata.get("playback_verified") is True
            or (isinstance(playback, dict) and playback.get("playing") is True)
        ):
            return True, {"method": details.get("method") or "media_playback_state", **details, "playback_verified": True}
        state_changed = bool(details.get("state_changed"))
        page_changed = bool(details.get("current_url") and details.get("current_url") != details.get("before_url")) or bool(details.get("current_title") and details.get("current_title") != details.get("before_title"))
        if state_changed or page_changed:
            return True, {"method": details.get("method") or "owned_tab_state_change", **details}
        # DOM click dispatch is evidence that the event reached the target, but not
        # proof that the application accepted it. Keep it successful-but-uncertain.
        return False, {"definitive_failure": False, "method": "dom_click_dispatch", "reason": "postcondition_not_observable", **details}

    # -------------------------------------------------------------
    # Browser Verification
    # -------------------------------------------------------------
    def _verify_search(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        tab_id = result.metadata.get("tab_id")
        query = action.parameters.get("query") or action.parameters.get("value")
        valid = bool(result.success and tab_id and query)
        return valid, {
            "method": "owned_tab_search_submission",
            "tab_id": tab_id,
            "query": query,
            "submission_observed": valid,
            "definitive_failure": not valid,
        }

    def _verify_navigate(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        """Verify navigation against the exact tab ADVI executed on."""
        expected_url = str(result.metadata.get("expected_url") or action.parameters.get("url") or "").strip()
        tab_id = result.metadata.get("tab_id")
        cdp_port = int(result.metadata.get("cdp_port", 9222))

        if not expected_url:
            return False, {"definitive_failure": False, "reason": "missing_expected_url"}
        if not tab_id:
            return False, {
                "definitive_failure": False,
                "reason": "missing_owned_tab_id",
            }

        current_url: str | None = None
        try:
            # Query the exact target by ID instead of selecting the first page in /json/list.
            from advi.capabilities.browser.cdp import CDPBrowser
            tab = CDPBrowser(port=cdp_port).get_tab_info(str(tab_id))
            if tab and str(tab.get("id")) == str(tab_id):
                current_url = tab.get("url")
        except Exception as exc:
            logger.debug("Exact CDP tab verification failed for %s: %s", tab_id, exc)

        if not current_url:
            return False, {
                "definitive_failure": False,
                "reason": "tab_url_unavailable",
                "tab_id": str(tab_id),
            }

        expected = expected_url.rstrip("/").lower()
        actual = str(current_url).rstrip("/").lower()
        matched = actual == expected
        if not matched:
            # Permit legitimate application redirects/normalization while still
            # requiring the same origin and the meaningful requested query data.
            # Gmail commonly rewrites /mail/?view=cm to /mail/u/0/?tf=cm and
            # normalizes URL encoding/empty parameters; that is a successful
            # navigation, not a verification failure.
            expected_parts = urllib.parse.urlsplit(expected)
            actual_parts = urllib.parse.urlsplit(actual)
            if (
                expected_parts.scheme == actual_parts.scheme
                and expected_parts.netloc == actual_parts.netloc
            ):
                expected_path = expected_parts.path.rstrip("/")
                actual_path = actual_parts.path.rstrip("/")
                if expected_parts.netloc == "mail.google.com" and expected_path.startswith("/mail") and actual_path.startswith("/mail"):
                    expected_q = urllib.parse.parse_qs(expected_parts.query, keep_blank_values=False)
                    actual_q = urllib.parse.parse_qs(actual_parts.query, keep_blank_values=False)
                    # Require each non-empty requested parameter to survive the
                    # Gmail redirect. Empty compose parameters such as su= are
                    # intentionally ignored. view=cm is represented by tf=cm.
                    query_ok = all(
                        any(str(value) == str(candidate) for candidate in actual_q.get(key, []))
                        for key, values in expected_q.items()
                        if key != "view"
                        for value in values
                        if str(value) != ""
                    )
                    if expected_q.get("view") == ["cm"]:
                        query_ok = actual_q.get("tf") == ["cm"] and query_ok
                    matched = query_ok
                else:
                    matched = (
                        expected_path == actual_path
                        and expected_parts.query == actual_parts.query
                    )

        return matched, {
            "expected_url": expected_url,
            "current_url": current_url,
            "tab_id": str(tab_id),
            "url_matched": matched,
            "definitive_failure": not matched,
            "source": "exact_cdp_tab",
        }

    def _verify_david_fallback(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        fallback = bool(result.metadata.get("fallback"))
        completed = bool(result.success and fallback)
        return completed, {
            "source": "david_execution_loop",
            "fallback": fallback,
            "completed": completed,
            "result_count": result.metadata.get("result_count", 0),
        }

    def _verify_contact_lookup(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        rows = result.data
        valid = isinstance(rows, list) and all(isinstance(row, dict) and row.get("email") for row in rows)
        # Zero matches is a valid, verified lookup outcome; the model needs to
        # distinguish "not found" from "lookup failed".
        if isinstance(rows, list) and not rows:
            return True, {"method": "trusted_contact_store", "count": 0, "match": result.metadata.get("match"), "valid_result": True}
        return valid, {"method": "trusted_contact_store", "count": len(rows) if isinstance(rows, list) else None, "match": result.metadata.get("match"), "valid_result": valid, "definitive_failure": not valid, "reason": "invalid_contact_result" if not valid else None}

    def _verify_memory_update(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = dict(result.verification_details or {})
        key = result.metadata.get("key") or details.get("key")
        category = result.metadata.get("category") or details.get("category")
        valid = bool(key and category and result.success)
        return valid, {"method": "sqlite_memory_persistence", "category": category, "key": key, "persisted": valid, "definitive_failure": not valid, "reason": "missing_persistence_identity" if not valid else None}

    def _verify_memory_forget(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        deleted_count = result.metadata.get("deleted_count")
        valid = bool(result.success and isinstance(deleted_count, int) and deleted_count > 0)
        return valid, {"method": "sqlite_memory_deletion", "deleted_count": deleted_count, "deleted": valid, "definitive_failure": not valid, "reason": "no_records_deleted" if not valid else None}

    def _verify_memory_retrieval(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        valid = isinstance(data, list)
        return valid, {"method": "memory_query", "valid_result": valid, "count": len(data) if valid else None}

    def _verify_research_web(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        text = result.data
        has_content = isinstance(text, str) and bool(text.strip())
        tab_id = result.metadata.get("tab_id")
        current_url = result.metadata.get("current_url")
        valid = has_content and bool(tab_id and current_url)
        return valid, {"method": "owned_browser_research", "tab_id": tab_id, "current_url": current_url, "content_available": has_content, "length": len(text) if isinstance(text, str) else None, "definitive_failure": not valid, "reason": "research_result_not_readable" if not valid else None}

    def _verify_email_read(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        valid = isinstance(data, list)
        return valid, {"method": "gmail_message_read", "valid_result": valid, "count": len(data) if valid else None}

    def _verify_email_draft_read(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        data = result.data
        draft_id = None
        if isinstance(data, dict):
            draft_id = data.get("draft_id") or data.get("id")
        else:
            draft_id = getattr(data, "draft_id", None)
        return bool(draft_id), {"method": "gmail_draft_identity", "draft_id": draft_id, "valid_result": bool(draft_id)}

    def _verify_email_draft_create(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        draft_id = result.metadata.get("draft_id") or (result.data.get("draft_id") if isinstance(result.data, dict) else None)
        return bool(draft_id), {"method": "gmail_draft_identity", "draft_id": draft_id}

    def _verify_email_draft_update(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        draft_id = result.metadata.get("draft_id") or getattr(result.data, "draft_id", None) or (result.data.get("id") if isinstance(result.data, dict) else None)
        return bool(draft_id), {"method": "gmail_draft_identity", "draft_id": draft_id}

    # -------------------------------------------------------------
    # Email Verification
    # -------------------------------------------------------------
    def _verify_email_send(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        return self._verify_email_message_result(result)

    def _verify_email_reply(self, action: ToolCall, result: ToolResult) -> tuple[bool, dict[str, Any]]:
        details = self._verify_email_message_result(result)[1]
        details["status"] = "replied" if details.get("message_id") else details.get("status")
        return bool(details.get("message_id")), details

    @staticmethod
    def _verify_email_message_result(result: ToolResult) -> tuple[bool, dict[str, Any]]:
        message_id = result.metadata.get("message_id") or (result.data.get("id") if isinstance(result.data, dict) else None)
        if message_id:
            return True, {"message_id": message_id, "status": "sent"}
        return False, {"reason": "missing_message_id"}
