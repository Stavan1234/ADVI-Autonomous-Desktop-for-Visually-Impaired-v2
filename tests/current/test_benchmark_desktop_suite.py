"""ADVI Desktop Suite: every task sets up and every check judges the end state correctly."""
from __future__ import annotations

import shutil
import zipfile

import pytest

from benchmarks.common.splits import stratified_split
from benchmarks.desktop.suite import _CHECKS, load_tasks, prepare, run_checks


def _task(tid):
    return next(t for t in load_tasks() if t["id"] == tid)


def test_every_task_is_well_formed_and_sets_up(tmp_path):
    tasks = load_tasks()
    assert len({t["id"] for t in tasks}) == len(tasks)
    for task in tasks:
        assert task["instruction"] and task["checks"] and task["domain"]
        for check in task["checks"]:
            assert next(iter(check)) in _CHECKS, (task["id"], check)
        folder = prepare(task, tmp_path)
        assert folder.is_dir()
        # Sandbox tasks name their own folder, so ADVI never works outside it.
        if task["domain"] in {"files", "notepad", "word"}:
            assert f"ADVI Bench\\{task['id']}" in task["instruction"], task["id"]


def test_capability_sweep_tasks_are_well_formed(tmp_path):
    from pathlib import Path

    tasks = load_tasks(Path("benchmarks/desktop/capabilities.json"))
    assert len({t["id"] for t in tasks}) == len(tasks) >= 30
    for task in tasks:
        for check in task["checks"]:
            assert next(iter(check)) in _CHECKS, (task["id"], check)
    # Sends and deletes are only ever checked up to the confirmation.
    for task in tasks:
        if task["domain"] == "email" and any("confirmation_matches" in c for c in task["checks"]):
            assert any("no_tool" in c for c in task["checks"]), task["id"]


def test_new_capability_checks(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    import docx

    book = Workbook()
    book.active.append(["Item", "Cost"])
    book.active.append(["Rent", 9000])
    book.save(tmp_path / "Budget.xlsx")
    assert _CHECKS["xlsx_has"]({"xlsx_has": "Budget.xlsx", "values": ["Rent", "9000"]}, tmp_path, {})[0]
    deck = Presentation()
    for _ in range(3):
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Solar power"
    deck.save(tmp_path / "Solar.pptx")
    assert _CHECKS["pptx_has"]({"pptx_has": "Solar.pptx", "min_slides": 3, "words": ["solar"]}, tmp_path, {})[0]
    document = docx.Document()
    document.add_heading("Habits", 0)
    document.add_paragraph("| not a table |")
    document.save(tmp_path / "Habits.docx")
    ok, detail = _CHECKS["docx_structure"]({"docx_structure": "Habits.docx", "min_headings": 1}, tmp_path, {})
    assert not ok and "markdown_left" in detail  # literal Markdown fails the check


def test_an_old_run_folder_is_cleared_through_the_recycle_bin(tmp_path):
    task = _task("advi-create-folder")
    folder = prepare(task, tmp_path)
    (folder / "Projects").mkdir()
    trashed = []
    prepare(task, tmp_path, trash=lambda p: trashed.append(p) or shutil.rmtree(p))
    assert trashed == [folder] and not (folder / "Projects").exists()


def test_archive_task_passes_only_when_done(tmp_path):
    task = _task("waa-archive-docx")
    folder = prepare(task, tmp_path)
    assert not all(c["passed"] for c in run_checks(task, folder, {}))
    (folder / "Archive").mkdir()
    for name in ("Doc01.docx", "Doc02.docx"):
        (folder / name).rename(folder / "Archive" / name)
    assert all(c["passed"] for c in run_checks(task, folder, {}))


def test_generated_count_file_has_exactly_the_stated_count(tmp_path):
    task = _task("waa-notepad-count")
    folder = prepare(task, tmp_path)
    assert (folder / "largefile.txt").read_text(encoding="utf-8").split().count("example") == 37
    (folder / "example_count.txt").write_text("37\n", encoding="utf-8")
    assert all(c["passed"] for c in run_checks(task, folder, {}))


def test_large_files_report_rejects_listing_everything(tmp_path):
    task = _task("waa-large-files-report")
    folder = prepare(task, tmp_path)
    (folder / "report.txt").write_text("example.txt\ntesting.bin\nphoto.jpg\n", encoding="utf-8")
    assert not run_checks(task, folder, {})[0]["passed"]
    (folder / "report.txt").write_text("testing.bin\n", encoding="utf-8")
    assert run_checks(task, folder, {})[0]["passed"]


def test_zip_and_hidden_checks(tmp_path):
    task = _task("waa-zip-backup")
    folder = prepare(task, tmp_path)
    with zipfile.ZipFile(folder / "Downloads" / "DownloadsBackup.zip", "w") as archive:
        archive.write(folder / "Downloads" / "SampleText.txt", "SampleText.txt")
    assert not run_checks(task, folder, {})[0]["passed"]  # notes.md is missing

    import ctypes
    hide = _task("waa-hide-file")
    folder = prepare(hide, tmp_path)
    assert not run_checks(hide, folder, {})[0]["passed"]
    ctypes.windll.kernel32.SetFileAttributesW(str(folder / "secret.txt"), 0x2)
    assert run_checks(hide, folder, {})[0]["passed"]


def test_answer_and_confirmation_checks():
    find = _task("advi-find-file")
    assert run_checks(find, None, {"answer": "It's in 2024\\March."})[0]["passed"]
    assert not run_checks(find, None, {"answer": "I couldn't find it."})[0]["passed"]
    whatsapp = _task("advi-whatsapp-confirm")
    record = {"confirmations": ["Before I do that, I need your confirmation: send the WhatsApp message "
                                "'hello from the benchmark' to New Chapter. Say yes to continue."], "tools": []}
    assert all(c["passed"] for c in run_checks(whatsapp, None, record))
    record["tools"] = [{"name": "whatsapp_send_message"}]
    assert not all(c["passed"] for c in run_checks(whatsapp, None, record))


def test_docx_table_needs_header_plus_three_rows(tmp_path):
    import docx

    task = _task("advi-word-table")
    folder = prepare(task, tmp_path)
    document = docx.Document()
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Station"
    document.save(folder / "Stations.docx")
    assert run_checks(task, folder, {})[0]["passed"]


def test_a_broken_check_fails_instead_of_crashing(tmp_path):
    task = {"id": "x", "checks": [{"text_equals": "nope.txt", "text": "a"}]}
    result = run_checks(task, tmp_path, {})[0]
    assert not result["passed"] and "FileNotFoundError" in result["detail"]


def test_split_is_stratified_and_repeatable():
    tasks = load_tasks()
    one = stratified_split(tasks, group_of=lambda t: t["domain"], id_of=lambda t: t["id"], dev_fraction=0.4, seed=1)
    two = stratified_split(tasks, group_of=lambda t: t["domain"], id_of=lambda t: t["id"], dev_fraction=0.4, seed=1)
    assert one == two and not set(one["dev"]) & set(one["test"])
    assert set(one["dev"]) | set(one["test"]) == {t["id"] for t in tasks}
    domains = {t["id"]: t["domain"] for t in tasks}
    assert {domains[t] for t in one["test"]} == set(domains.values())  # every domain is scored


def test_runner_scripts_compile():
    import py_compile
    from pathlib import Path

    for path in Path("benchmarks").rglob("*.py"):
        py_compile.compile(str(path), doraise=True)
