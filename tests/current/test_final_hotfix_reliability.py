from pathlib import Path
from types import SimpleNamespace

from advi.tools.contracts import ToolCall, ToolResult
from advi.core.verification import ToolVerifier


def test_structured_file_content_mismatch_is_a_definitive_failure(tmp_path):
    import fitz
    path = tmp_path / "demo.pdf"
    doc = fitz.open(); page = doc.new_page(); page.insert_text((50, 80), "Doreamon") ; doc.save(path); doc.close()
    result = ToolResult(call_id="1", name="create_file", success=True, output=str(path), metadata={"resolved_path":str(path)})
    checked = ToolVerifier().verify(ToolCall("1", "create_file", {"path":str(path),"content":"A very different source request"}), result)
    assert checked.success is False
    assert checked.verified is False
    assert checked.verification_details["definitive_failure"] is True
    assert checked.verification_details["content_matches"] is False


def test_packaged_whatsapp_window_rejects_word_document():
    from advi.capabilities.desktop.executor import DesktopCapability
    cap = DesktopCapability()
    class Record:
        runtime_type="packaged"
        executable=None
        command=r"shell:AppsFolder\\5319275A.WhatsAppDesktop_cv1g1gvanyjgm!App"
    cap._window_title=lambda hwnd: "WhatsApp Case Study - Word (Product Activation Failed)"
    cap._window_process_id=lambda hwnd: 42
    cap._process_name=lambda pid: "WINWORD.EXE"
    assert cap._application_window_matches_record("WhatsApp", 100, Record()) is False


def test_system_capability_returns_local_status():
    from advi.capabilities.system import SystemCapability
    result = SystemCapability().execute(ToolCall("s1", "get_system_info", {}))
    assert result.success and result.verified
    assert "os" in result.output and "ram_total_gb" in result.output


def test_search_intent_preserves_exact_literal():
    from advi.fallback.intent_service import IntentService
    class P:
        calls=[]
        def chat(self, *args, **kwargs): raise AssertionError("provider should not be used")
    intent=IntentService(P(), "src/advi/fallback/prompts/intent_prompt.txt").extract_intent("open google chrome and search for youtube")
    assert intent.goal == "web_search"
    assert intent.entities["query"] == "youtube"


def test_planner_web_search_preserves_exact_literal():
    from advi.fallback.planner_service import PlannerService
    from advi.fallback.intent_schema import Intent
    class P: pass
    planner=PlannerService(P(), "src/advi/fallback/prompts/planner_prompt.txt")
    plan=planner._build_deterministic_plan(Intent(goal="web_search", entities={"query":"youtube"}))
    assert plan is not None
    assert "q=youtube" in plan.actions[1].parameters["url"]



def test_execution_loop_refuses_to_send_if_prepared_gmail_fields_changed():
    from advi.fallback.execution_loop import ExecutionLoop
    from advi.fallback.execution_context import ExecutionContext

    calls = []
    class FakeCDP:
        def verify_gmail_compose(self, *args, **kwargs):
            calls.append(("verify", kwargs))
            return False
        def send_gmail_compose(self, *args, **kwargs):
            calls.append(("send", kwargs))
            return True
    class FakeExecutor:
        def __init__(self):
            self.cdp_browser = FakeCDP()

    context = ExecutionContext()
    context.target_tab_id = "tab-1"
    loop = ExecutionLoop(executor=FakeExecutor(), context=context)
    result = loop.send_prepared_email(recipient="person@example.com", subject="Expected", body="Expected body")
    assert result.success is False
    assert result.metadata["reason"] == "prepared_gmail_fields_changed"
    assert not any(kind == "send" for kind, _ in calls)


def test_david_focus_uses_kernel32_get_current_thread_id():
    import inspect
    from advi.fallback import executor
    source=inspect.getsource(executor.FallbackExecutor._focus_hwnd)
    assert "kernel32.GetCurrentThreadId" in source
    assert "user32.GetCurrentThreadId" not in source


def test_execution_loop_staged_email_is_never_a_send_in_stage_one():
    from advi.fallback.execution_loop import ExecutionLoop, ExecutionLoopResult
    from advi.fallback.execution_context import ExecutionContext
    from advi.fallback.action_plan_schema import Action

    calls=[]
    class FakeCDP:
        def verify_gmail_compose(self, *args, **kwargs): calls.append("verify"); return True
        def send_gmail_compose(self, *args, **kwargs): calls.append("send"); return True
    class FakeExecutor:
        def __init__(self): self.cdp_browser=FakeCDP()
        def execute_action(self, action, context=None):
            calls.append(action.action)
            if action.action == "open_application":
                context.current_application="Chrome"
                context.target_tab_id="tab-1"
            elif action.action == "navigate":
                context.target_tab_id="tab-1"
            return SimpleNamespace(action=action.action, success=True, error=None, data=action.parameters.get("url"))
    loop=ExecutionLoop(executor=FakeExecutor(), context=ExecutionContext())
    result=loop.prepare_email_draft(recipient="person@example.com", subject="Test", body="Body")
    assert result.success and result.metadata["awaiting_send_confirmation"]
    assert "send" not in calls
    assert calls[-1] == "verify"


def test_pptx_renderer_creates_native_title_and_table(tmp_path):
    from advi.capabilities.files.executor import FileCapability
    from advi.tools.contracts import ToolCall
    from pptx import Presentation
    cap=FileCapability(); cap.resolve_path=lambda raw: tmp_path / str(raw)
    content="Doreamon\nCharacters\n| Character | Role |\n| --- | --- |\n| Doraemon | Robotic cat |"
    result=cap.execute(ToolCall("p","create_file",{"path":"styled.pptx","content":content,"slides":2,"style":"polished-colorful"}))
    assert result.success
    presentation=Presentation(tmp_path/"styled.pptx")
    assert len(presentation.slides)==2
    assert presentation.slides[0].shapes.title.text == "Doreamon"
    assert any(getattr(shape, "has_table", False) for slide in presentation.slides for shape in slide.shapes)
