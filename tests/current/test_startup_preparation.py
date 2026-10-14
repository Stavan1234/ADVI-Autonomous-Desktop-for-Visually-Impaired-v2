from __future__ import annotations

import sys
import types


def test_memory_retriever_preload_happens_once(monkeypatch, tmp_path):
    from advi.memory.long_term import LongTermMemory
    from advi.memory.retriever import MemoryRetriever

    class FakeModel:
        instances = 0

        def __init__(self, name):
            FakeModel.instances += 1
            self.name = name

    fake_module = types.ModuleType("sentence_transformers")
    fake_module.SentenceTransformer = FakeModel
    monkeypatch.setitem(sys.modules, "sentence_transformers", fake_module)

    store = LongTermMemory(tmp_path / "memory.db")
    retriever = MemoryRetriever(store)

    assert retriever.semantic_ready is False
    assert retriever.preload() is True
    assert retriever.preload() is True
    assert retriever.semantic_ready is True
    assert FakeModel.instances == 1



def test_memory_retriever_does_not_retry_failed_semantic_startup(monkeypatch, tmp_path):
    from advi.memory.long_term import LongTermMemory
    from advi.memory.retriever import MemoryRetriever

    calls = []

    class FailingModel:
        def __init__(self, name):
            calls.append(name)
            raise RuntimeError("download failed")

    fake_module = types.ModuleType("sentence_transformers")
    fake_module.SentenceTransformer = FailingModel
    monkeypatch.setitem(sys.modules, "sentence_transformers", fake_module)

    store = LongTermMemory(tmp_path / "memory.db")
    retriever = MemoryRetriever(store)

    assert retriever.preload() is False
    assert retriever.preload() is False
    assert calls == ["BAAI/bge-small-en-v1.5"]
    assert retriever.semantic_error == "download failed"



def test_provider_session_manager_prepare_creates_session_without_user_request():
    from advi.agent.provider_session import ProviderSessionManager
    from advi.providers.base import LLMProvider

    class FakeProvider(LLMProvider):
        @property
        def name(self):
            return "fake"

        @property
        def model(self):
            return "fake-model"

        def chat(self, messages, *, purpose="chat"):
            raise AssertionError("prepare must not send a model request")

    manager = ProviderSessionManager([FakeProvider()])
    session = manager.prepare(
        system_instruction="stable instructions",
        tools=[],
    )

    assert manager.session is session
    assert session.messages == [{"role": "system", "content": "stable instructions"}]


def test_piper_preflight_only_checks_assets(tmp_path):
    from advi.io.tts import PiperTTS

    exe = tmp_path / "piper.exe"
    model = tmp_path / "voice.onnx"
    config = tmp_path / "voice.onnx.json"
    output = tmp_path / "audio" / "out.wav"
    exe.write_bytes(b"")
    model.write_bytes(b"")
    config.write_text("{}", encoding="utf-8")

    tts = PiperTTS(exe, model, output)
    tts.synthesize = lambda _text: True

    assert tts.preflight() is True
    assert output.parent.is_dir()
    assert not output.exists()


def test_gmail_freeze_prevents_hidden_late_initialization(monkeypatch, tmp_path):
    import advi.integrations.gmail.service as gmail_module

    monkeypatch.setattr(gmail_module, "TOKEN_FILE", tmp_path / "missing-token.json")
    monkeypatch.setattr(gmail_module, "CREDENTIALS_FILE", tmp_path / "credentials.json")

    service = gmail_module.LazyGmailService()
    assert service.prepare() is False
    service.freeze_lazy_initialization()

    try:
        _ = service.users
    except RuntimeError as exc:
        assert "not initialized" in str(exc).lower()
    else:
        raise AssertionError("Gmail attempted hidden lazy initialization after startup freeze")

def test_gmail_first_user_action_can_start_oauth_after_startup_freeze(monkeypatch, tmp_path):
    import advi.integrations.gmail.service as gmail_module

    monkeypatch.setattr(gmail_module, "TOKEN_FILE", tmp_path / "missing-token.json")
    monkeypatch.setattr(gmail_module, "CREDENTIALS_FILE", tmp_path / "credentials.json")

    calls = []

    class FakeGmailService:
        def __init__(self):
            calls.append("initialize")

    monkeypatch.setattr(gmail_module, "GmailService", FakeGmailService)

    service = gmail_module.LazyGmailService()
    service.freeze_lazy_initialization()

    assert service.ensure_ready() is True
    assert service.initialized is True
    assert calls == ["initialize"]


def test_gmail_capability_explicitly_initializes_lazy_service(monkeypatch):
    from advi.capabilities.gmail.executor import GmailCapability
    from advi.tools.contracts import ToolCall

    class FakeLazyService:
        def __init__(self):
            self.ready_calls = 0

        def ensure_ready(self):
            self.ready_calls += 1

        def health_check(self):
            return True

        def list_messages(self, query="", max_results=10):
            return []

    service = FakeLazyService()
    capability = GmailCapability(service)
    result = capability.execute(ToolCall("g1", "email_read", {"query": "", "limit": 1}))

    assert service.ready_calls == 1
    assert result.success is True



def test_gmail_first_use_invokes_installed_app_oauth_browser(monkeypatch, tmp_path):
    import sys
    import types
    import advi.integrations.gmail.auth as auth_module

    credentials_file = tmp_path / "gmail_credentials.json"
    token_file = tmp_path / "gmail_token.json"
    credentials_file.write_text('{"installed":{"client_id":"test"}}', encoding="utf-8")
    monkeypatch.setattr(auth_module, "CREDENTIALS_FILE", credentials_file)
    monkeypatch.setattr(auth_module, "TOKEN_FILE", token_file)

    calls = []

    class FakeCredentials:
        valid = False
        expired = False
        refresh_token = None

        def to_json(self):
            return '{"access_token":"fresh"}'

    class FakeFlow:
        @classmethod
        def from_client_secrets_file(cls, filename, scopes):
            calls.append(("from_client_secrets_file", filename, list(scopes)))
            return cls()

        def run_local_server(self, *, port, open_browser):
            calls.append(("run_local_server", port, open_browser))
            return FakeCredentials()

    google = types.ModuleType("google")
    google_auth = types.ModuleType("google.auth")
    google_auth_transport = types.ModuleType("google.auth.transport")
    google_auth_requests = types.ModuleType("google.auth.transport.requests")
    google_oauth2 = types.ModuleType("google.oauth2")
    google_credentials = types.ModuleType("google.oauth2.credentials")
    google_flow_pkg = types.ModuleType("google_auth_oauthlib")
    google_flow = types.ModuleType("google_auth_oauthlib.flow")

    google_auth_requests.Request = type("Request", (), {})
    google_credentials.Credentials = type("Credentials", (), {
        "from_authorized_user_file": staticmethod(lambda *args, **kwargs: None),
    })
    google_flow.InstalledAppFlow = FakeFlow

    monkeypatch.setitem(sys.modules, "google", google)
    monkeypatch.setitem(sys.modules, "google.auth", google_auth)
    monkeypatch.setitem(sys.modules, "google.auth.transport", google_auth_transport)
    monkeypatch.setitem(sys.modules, "google.auth.transport.requests", google_auth_requests)
    monkeypatch.setitem(sys.modules, "google.oauth2", google_oauth2)
    monkeypatch.setitem(sys.modules, "google.oauth2.credentials", google_credentials)
    monkeypatch.setitem(sys.modules, "google_auth_oauthlib", google_flow_pkg)
    monkeypatch.setitem(sys.modules, "google_auth_oauthlib.flow", google_flow)

    creds = auth_module.get_gmail_credentials()

    assert isinstance(creds, FakeCredentials)
    assert calls[0][0] == "from_client_secrets_file"
    assert calls[1] == ("run_local_server", 0, True)
    assert token_file.exists()
    assert "fresh" in token_file.read_text(encoding="utf-8")


def test_memory_search_does_not_start_a_background_model_loader(monkeypatch, tmp_path):
    from advi.memory.long_term import Memory
    from advi.memory.long_term import LongTermMemory
    from advi.memory.retriever import MemoryRetriever

    store = LongTermMemory(tmp_path / "memory.db")
    memory = Memory(
        id=1,
        category="user",
        key="name",
        value="Stavan",
        statement="The user's name is Stavan.",
        confidence=1.0,
        created_at="2026-01-01",
        updated_at="2026-01-01T00:00:00",
    )
    store._memories_for_tests = [memory]
    monkeypatch.setattr(store, "search", lambda query, limit=15: [memory])
    monkeypatch.setattr(store, "recent", lambda limit=32: [memory])

    retriever = MemoryRetriever(store)
    results = retriever.search("what is my name", limit=1)

    assert results
    assert results[0].memory.statement == "The user's name is Stavan."
    assert results[0].semantic_score == 0.0


def test_semantic_preload_swallows_model_console_output(monkeypatch, tmp_path, capsys):
    import sys
    import types

    from advi.memory.long_term import LongTermMemory
    from advi.memory.retriever import MemoryRetriever

    class NoisyModel:
        def __init__(self, name):
            print("Loading weights: 100%")
            print("Warning: unauthenticated requests to the HF Hub", file=sys.stderr)

    fake_module = types.ModuleType("sentence_transformers")
    fake_module.SentenceTransformer = NoisyModel
    monkeypatch.setitem(sys.modules, "sentence_transformers", fake_module)

    retriever = MemoryRetriever(LongTermMemory(tmp_path / "memory.db"))
    assert retriever.preload() is True
    captured = capsys.readouterr()
    assert "Loading weights" not in captured.out
    assert "HF Hub" not in captured.err
    assert retriever.semantic_ready is True


def test_startup_announcer_speaks_first_message_before_background_work(capsys):
    from advi.io.console import StartupAnnouncer

    spoken = []
    announcer = StartupAnnouncer(interval_seconds=30, speaker=lambda text: spoken.append(text) or True)
    announcer.start()
    try:
        assert spoken == ["Just preparing everything, sir..."]
        captured = capsys.readouterr()
        assert "Just preparing everything, sir..." in captured.out
    finally:
        announcer.stop()


def test_semantic_preload_is_quiet(monkeypatch, tmp_path, caplog):
    import sys
    import types
    from advi.memory.long_term import LongTermMemory
    from advi.memory.retriever import MemoryRetriever

    class NoisyModel:
        def __init__(self, name):
            import logging
            logging.getLogger("huggingface_hub").warning("NOISY SHOULD BE HIDDEN")
            print("NOISY STDOUT")
            print("NOISY STDERR", file=sys.stderr)

    fake_module = types.ModuleType("sentence_transformers")
    fake_module.SentenceTransformer = NoisyModel
    monkeypatch.setitem(sys.modules, "sentence_transformers", fake_module)

    caplog.set_level("WARNING")
    retriever = MemoryRetriever(LongTermMemory(tmp_path / "memory.db"))
    assert retriever.preload() is True
    assert "NOISY SHOULD BE HIDDEN" not in caplog.text


def test_create_file_generates_real_common_document_containers(tmp_path, monkeypatch):
    from advi.capabilities.files.executor import FileCapability
    from advi.tools.contracts import ToolCall

    capability = FileCapability()
    monkeypatch.setattr(capability, "resolve_path", lambda raw: tmp_path / str(raw))

    for filename, content in (
        ("example.docx", "Hello from ADVI"),
        ("example.xlsx", "Name,Value\nADVI,1"),
        ("example.pptx", "Title\nBody"),
        ("example.pdf", "Hello from ADVI"),
    ):
        result = capability.execute(ToolCall(filename, "create_file", {"path": filename, "content": content}))
        assert result.success, (filename, result.error)
        path = tmp_path / filename
        assert path.is_file() and path.stat().st_size > 100

    from docx import Document
    assert "Hello from ADVI" in [p.text for p in Document(tmp_path / "example.docx").paragraphs]

    from openpyxl import load_workbook
    workbook = load_workbook(tmp_path / "example.xlsx", read_only=True, data_only=True)
    try:
        assert workbook.active["A1"].value == "Name"
        assert workbook.active["B2"].value == 1
    finally:
        workbook.close()

    from pptx import Presentation
    presentation = Presentation(tmp_path / "example.pptx")
    assert presentation.slides[0].shapes.title.text == "Title"

    import fitz
    pdf = fitz.open(tmp_path / "example.pdf")
    try:
        assert "Hello from ADVI" in pdf[0].get_text()
    finally:
        pdf.close()


def test_save_file_generates_real_structured_document_containers(tmp_path, monkeypatch):
    from advi.capabilities.files.executor import FileCapability
    from advi.tools.contracts import ToolCall
    import fitz
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    capability = FileCapability()
    monkeypatch.setattr(capability, "resolve_path", lambda raw: tmp_path / str(raw))

    pdf_result = capability.execute(ToolCall("pdf", "save_file", {"path": "report.pdf", "content": "India report — 2026 BC"}))
    assert pdf_result.success
    pdf = fitz.open(tmp_path / "report.pdf")
    try:
        assert "India report" in pdf[0].get_text()
    finally:
        pdf.close()

    docx_result = capability.execute(ToolCall("docx", "save_file", {"path": "report.docx", "content": "India report"}))
    assert docx_result.success
    assert "India report" in [p.text for p in Document(tmp_path / "report.docx").paragraphs]

    xlsx_result = capability.execute(ToolCall("xlsx", "save_file", {"path": "report.xlsx", "content": "Name,Value\nIndia,1"}))
    assert xlsx_result.success
    workbook = load_workbook(tmp_path / "report.xlsx", read_only=True, data_only=True)
    try:
        assert workbook.active["A2"].value == "India"
        assert workbook.active["B2"].value == 1
    finally:
        workbook.close()

    pptx_result = capability.execute(ToolCall("pptx", "save_file", {"path": "report.pptx", "content": "Title\nBody"}))
    assert pptx_result.success
    presentation = Presentation(tmp_path / "report.pptx")
    assert presentation.slides[0].shapes.title.text == "Title"


def test_pdf_verifier_accepts_native_structured_content(tmp_path, monkeypatch):
    from advi.capabilities.files.executor import FileCapability
    from advi.core.verification import ToolVerifier
    from advi.tools.contracts import ToolCall, ToolResult

    capability = FileCapability()
    monkeypatch.setattr(capability, "resolve_path", lambda raw: tmp_path / str(raw))
    content = "India report — 2026 BC"
    result = capability.execute(ToolCall("pdf", "create_file", {"path": "report.pdf", "content": content}))
    assert result.success

    verified, details = ToolVerifier()._verify_create_file(
        ToolCall("pdf", "create_file", {"path": "report.pdf", "content": content}),
        result,
    )
    assert verified, details
